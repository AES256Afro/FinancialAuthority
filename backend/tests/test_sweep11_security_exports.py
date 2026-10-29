"""Security sweep 11, part three: every export format opened back, and the restore drill's temporary copies.

A second person with the "everything" role sees what the owner sees, by design (docs/SECURITY.md, Who can reach
it), so their export is the owner's; a "shared" person reaches no export at all, and a key none. What is checked
here is that each format opened back holds the rows of the accounts the maker may see, and nothing of a person the
maker is not; and that a drill leaves nothing behind."""

from __future__ import annotations

import csv
import io
import json
import tempfile
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.db import engine
from app.engine.household import seed
from app.models import Account, Transaction

PW = "correct horse battery"


def _invite(owner: TestClient, username: str, role: str) -> str:
    r = owner.post("/api/v1/people", json={"username": username, "display_name": username.title(), "role": role})
    assert r.status_code == 200, r.text
    return r.json()["invite"]["link"].split("#", 1)[1]


def _join(app, token: str, password: str = "a long enough passphrase") -> TestClient:  # noqa: ANN001
    person = TestClient(app)
    person.headers["X-Requested-With"] = "FinancialAuthority"
    r = person.post("/api/auth/join", json={"token": token, "password": password})
    assert r.status_code == 200, r.text
    return person


def _open_back(fmt: str, content: bytes) -> str:
    """The text of an export as a reader would see it."""
    if fmt == "csv":
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            return "\n".join(z.read(n).decode("utf-8", "replace") for n in z.namelist())
    if fmt == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True)
        return "\n".join(",".join("" if c is None else str(c) for c in row) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
    if fmt == "pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    return content.decode("utf-8", "replace")


def test_every_format_opened_back_holds_only_what_its_maker_may_see(owner: TestClient, client: TestClient):
    with Session(engine) as session:
        seed(session)
    alex = _join(client.app, _invite(owner, "alex", "everything"))
    sam = _join(client.app, _invite(owner, "sam", "shared"))
    sam_id = next(p["id"] for p in owner.get("/api/v1/people").json()["people"] if p["username"] == "sam")
    alex_id = next(p["id"] for p in owner.get("/api/v1/people").json()["people"] if p["username"] == "alex")
    with Session(engine) as session:
        # an account of Sam's own, and one of Alex's own, each with a transaction nobody else should see as theirs
        for pid, name, memo in ((sam_id, "Sam private checking", "SAMONLY-ROW-7731"), (alex_id, "Alex private savings", "ALEXONLY-ROW-9912")):
            account = Account(name=name, institution="Example Bank", kind="checking", connector="manual", external_id=f"p{pid}", person_id=pid, shared=False)
            session.add(account)
            session.flush()
            from datetime import date

            session.add(Transaction(account_id=account.id, date=date(2026, 10, 10), amount=12.34, raw_name=memo, merchant=memo, description=memo, category="Other", source="manual", row_hash=f"sweep11-{pid}"))
        session.commit()
    for fmt in ("csv", "xlsx", "json", "pdf", "ics"):
        mine = owner.get("/api/v1/exports/file", params={"format": fmt})
        assert mine.status_code == 200, (fmt, mine.text)
        theirs = alex.get("/api/v1/exports/file", params={"format": fmt})
        assert theirs.status_code == 200, (fmt, theirs.text)  # "everything" sees what the owner sees
        assert sam.get("/api/v1/exports/file", params={"format": fmt}).status_code == 403, fmt  # "shared" exports nothing
        assert TestClient(client.app).get("/api/v1/exports/file", params={"format": fmt}).status_code == 401
        text = _open_back(fmt, theirs.content)
        if fmt in ("csv", "json", "xlsx"):
            assert "SAMONLY-ROW-7731" in text and "ALEXONLY-ROW-9912" in text, fmt  # the everything role is the owner's view, by design
        assert "a long enough passphrase" not in text and "scrypt$" not in text, fmt
    # the calendar feed a person makes: the shared role cannot make one; the everything role's is the household's
    assert sam.post("/api/v1/calendar/feed", json={"amounts": False}).status_code == 403
    assert alex.post("/api/v1/calendar/feed", json={"amounts": False}).status_code == 200


def test_the_restore_drill_removes_its_temporary_copies(owner: TestClient, tmp_path: Path, monkeypatch):
    from app import offbox, restore
    from app.config import settings

    monkeypatch.setattr(settings, "restore_dir", tmp_path / "points")
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path / "tmp"))
    (tmp_path / "tmp").mkdir()
    with Session(engine) as session:
        seed(session)
    restore.make_point("daily", "sweep 11")
    before = sorted(p.name for p in (tmp_path / "tmp").iterdir())
    with Session(engine) as session:
        result = offbox.drill(session)
    assert result["ok"], result
    after = sorted(p.name for p in (tmp_path / "tmp").iterdir())
    assert after == before, f"the drill left {set(after) - set(before)} behind"
    kept = sorted(p.name for p in (tmp_path / "points").iterdir())
    assert all("daily" in p and "drill" not in p for p in kept), kept  # the point and its sidecars; no copy of the drill's
