"""Documents the owner keeps: a lease, its addenda, the property's charge sheets, a utility's terms, a card's
agreement, a receipt. Each is kept with what it belongs to, the file encrypted with the data key under the data
folder, its text read page by page (a PDF's text layer, a spreadsheet's cells, a text file) and kept encrypted.

The terms a lease states (rent, fees, the due day, the lease's end, the notice it needs, the late fee, the deposit)
are offered as a fact on Needs you; nothing changes until the owner confirms. Models read documents only through
the tools, which return the text boxed as untrusted data; nothing here trains a model."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

from sqlmodel import Session, select

from .. import crypto
from ..config import settings
from ..models import Document, ExtractedFact

MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 60
MAX_PAGE_CHARS = 20_000
XLSX_MAX_EXPANDED = 16 * 1024 * 1024  # what a spreadsheet's parts may add up to, unzipped
XLSX_MAX_ROWS = 5_000
XLSX_MAX_COLS = 60
KINDS = ("lease", "addendum", "charges", "welcome", "utility_terms", "card_agreement", "receipt", "statement", "paystub", "other")
TYPES = {"pdf": "application/pdf", "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "txt": "text/plain", "csv": "text/csv", "md": "text/plain",
         "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
ATTACH = re.compile(r"^(?:home|utility|bill|account|transaction|debt|paystub):\d+$")


def folder() -> Path:
    path = Path(settings.data_dir) / "documents"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _seal(raw: bytes) -> bytes:
    return crypto.encrypt(base64.b64encode(raw).decode()).encode() if settings.data_key else raw


def _open(stored: bytes) -> bytes:
    text = stored.decode(errors="ignore")
    return base64.b64decode(crypto.decrypt(text)) if crypto.is_encrypted(text) else stored


def read_pages(raw: bytes, ext: str) -> list[str]:
    """The text of each page (a sheet counts as a page)."""
    return read_pages_with(raw, ext)[0]


def read_pages_with(raw: bytes, ext: str) -> tuple[list[str], str]:
    """The text of each page, and how it was had: "" from the file's own text, "ocr" read from a photo, or from a
    scanned PDF whose pages are pictures, by the free reader on this server (app/ocr.py) when it is installed."""
    from .. import ocr

    if ext in ("png", "jpg", "jpeg"):
        text = ocr.image_text(raw)
        return ([text], "ocr") if text else ([], "")
    pages: list[str] = []
    try:
        if ext == "pdf":
            from ..parsers.paystub import pdf_pages

            pages = pdf_pages(raw, max_pages=MAX_PAGES, seconds=10.0, page_chars=MAX_PAGE_CHARS)
        elif ext == "xlsx":
            import zipfile

            from openpyxl import load_workbook

            # A spreadsheet is a zip: what it expands to is checked before it is opened, and rows are counted
            # whether or not they hold anything, so a tiny file can neither fill the memory nor loop.
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                expanded = sum(i.file_size for i in z.infolist())
            if expanded > XLSX_MAX_EXPANDED:
                return [], ""
            wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            for ws in list(wb.worksheets)[:MAX_PAGES]:
                lines, size = [ws.title], 0
                for n, row in enumerate(ws.iter_rows(values_only=True)):
                    if n >= XLSX_MAX_ROWS or size > MAX_PAGE_CHARS:
                        break
                    cells = [str(v) for v in row[:XLSX_MAX_COLS] if v not in (None, "")]
                    if cells:
                        line = "  ".join(cells)
                        lines.append(line)
                        size += len(line)
                pages.append("\n".join(lines)[:MAX_PAGE_CHARS])
        elif ext in ("txt", "csv", "md"):
            text = raw.decode("utf-8-sig", errors="replace")
            pages = [text[i:i + MAX_PAGE_CHARS] for i in range(0, min(len(text), MAX_PAGE_CHARS * MAX_PAGES), MAX_PAGE_CHARS)] or [""]
    except Exception:  # noqa: BLE001 - a file that cannot be read is kept, with no text
        pages = []
    if ext == "pdf" and not any(p.strip() for p in pages):
        scanned = ocr.scan_pages(raw)
        if any(p.strip() for p in scanned):
            return [p[:MAX_PAGE_CHARS] for p in scanned], "ocr"
    return pages, ""


def save_document(session: Session, raw: bytes, *, filename: str, title: str = "", kind: str = "other", attach_to: str = "", note: str = "", period: str = "", merchant: str = "",
                  amount: float | None = None, bought_on=None, return_by=None, warranty_until=None, source: str = "upload") -> Document:  # noqa: ANN001 - dates
    if len(raw) > MAX_BYTES:
        raise ValueError("A document is at most 25 MB")
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext not in TYPES:
        raise ValueError("A PDF, a photo (PNG or JPEG), a spreadsheet (XLSX) or a text file")
    if ext == "pdf" and not raw.startswith(b"%PDF"):
        raise ValueError("That file is not a PDF")
    if kind not in KINDS:
        kind = "other"
    if attach_to and not ATTACH.fullmatch(attach_to):
        raise ValueError("attach_to names what it belongs to, like home:3")
    digest = hashlib.sha256(raw).hexdigest()
    path = folder() / f"{digest}.bin"
    if not path.exists():
        path.write_bytes(_seal(raw))
        path.chmod(0o600)  # sealed, and the app's user only
    pages, read_with = read_pages_with(raw, ext)
    if period and not re.fullmatch(r"\d{4}-\d{2}", period):
        raise ValueError("A month is written YYYY-MM")
    if not period and bought_on is not None:
        period = bought_on.isoformat()[:7]
    doc = Document(title=(title or filename).strip()[:120], kind=kind, attach_to=attach_to, filename=Path(filename).name[:120], content_type=TYPES[ext], size=len(raw), sha256=digest,
                   pages_enc=crypto.encrypt(json.dumps(pages)) if settings.data_key else json.dumps(pages), pages=len(pages), note=note.strip()[:300],
                   period=period, merchant=merchant.strip()[:80], amount=round(float(amount), 2) if amount is not None else None, bought_on=bought_on, return_by=return_by,
                   warranty_until=warranty_until, source=source[:60], read_with=read_with)
    session.add(doc)
    session.commit()
    session.refresh(doc)
    if kind in ("lease", "addendum", "charges", "welcome") and attach_to.startswith("home:"):
        offer_terms(session, doc, pages)
    return doc


def pages_of(doc: Document) -> list[str]:
    raw = doc.pages_enc or "[]"
    try:
        return json.loads(crypto.decrypt(raw) if crypto.is_encrypted(raw) else raw)
    except Exception:  # noqa: BLE001 - unreadable without the key that sealed it
        return []


def file_of(doc: Document) -> bytes:
    path = folder() / f"{doc.sha256}.bin"
    if not path.exists():
        raise FileNotFoundError(doc.sha256)
    return _open(path.read_bytes())


def remove_document(session: Session, doc: Document) -> None:
    """The row goes now; the file goes only when no other row and no restore point names it, so a restore can
    bring the document back."""
    from .. import restore

    sha = doc.sha256
    session.delete(doc)
    session.commit()
    restore.release_document_file(sha)


def suggested(doc: Document, today: Any = None) -> dict[str, Any] | None:
    """For a receipt kept without its total or day: what its words seem to say, offered on the screen to keep or
    not. Never applied on its own."""
    if doc.kind != "receipt" or (doc.amount is not None and doc.bought_on is not None):
        return None
    from datetime import date as _date

    from ..ocr import receipt_reading

    got = receipt_reading("\n".join(pages_of(doc)), today or _date.today())
    out = {k: v for k, v in (("amount", got.get("total")), ("bought_on", got.get("bought_on"))) if v is not None and getattr(doc, k) is None}
    return out or None


def doc_out(doc: Document) -> dict[str, Any]:
    return {"read_with": doc.read_with, "suggested": suggested(doc),"id": doc.id, "title": doc.title, "kind": doc.kind, "attach_to": doc.attach_to, "filename": doc.filename, "content_type": doc.content_type, "size": doc.size,
            "pages": doc.pages, "has_text": doc.pages > 0 and any(p.strip() for p in pages_of(doc)), "note": doc.note, "uploaded_at": doc.uploaded_at.isoformat(), "source": f"document:{doc.id}",
            "period": doc.period, "merchant": doc.merchant, "amount": doc.amount, "bought_on": doc.bought_on.isoformat() if doc.bought_on else None,
            "return_by": doc.return_by.isoformat() if doc.return_by else None, "warranty_until": doc.warranty_until.isoformat() if doc.warranty_until else None, "came_from": doc.source}


def save_text(session: Session, text: str, *, filename: str, title: str, kind: str, attach_to: str = "", note: str = "", **facts: Any) -> Document:
    """A document made from text the app already holds (a receipt email's words), kept like an upload."""
    return save_document(session, text.encode("utf-8"), filename=filename if filename.endswith(".txt") else filename + ".txt", title=title, kind=kind, attach_to=attach_to, note=note, **facts)


def update_document(session: Session, doc: Document, values: dict[str, Any]) -> Document:
    """What the owner changes on a kept file: its title, kind, what it belongs to, its month and a receipt's facts."""
    from datetime import date as _date

    if "attach_to" in values:
        attach = str(values["attach_to"] or "")
        if attach and not ATTACH.fullmatch(attach):
            raise ValueError("attach_to names what it belongs to, like transaction:12")
        doc.attach_to = attach
    if "period" in values:
        period = str(values["period"] or "")
        if period and not re.fullmatch(r"\d{4}-\d{2}", period):
            raise ValueError("A month is written YYYY-MM")
        doc.period = period
    for key in ("title", "note", "merchant"):
        if values.get(key) is not None:
            setattr(doc, key, str(values[key]).strip()[:120 if key != "note" else 300])
    if values.get("kind") is not None:
        doc.kind = values["kind"] if values["kind"] in KINDS else doc.kind
    if "amount" in values:
        doc.amount = round(float(values["amount"]), 2) if values["amount"] is not None else None
    for key in ("bought_on", "return_by", "warranty_until"):
        if key in values:
            v = values[key]
            setattr(doc, key, _date.fromisoformat(v) if isinstance(v, str) and v else v or None)
    if doc.kind == "receipt" and not doc.period and doc.bought_on:
        doc.period = doc.bought_on.isoformat()[:7]
    session.add(doc)
    session.commit()
    session.refresh(doc)
    return doc


def search(session: Session, query: str, limit: int = 10) -> list[dict[str, Any]]:
    """Each place the words appear, with the page and a few lines around them."""
    words = [w for w in re.findall(r"[\w$.%-]+", query.lower()) if len(w) > 1][:8]
    if not words:
        return []
    hits = []
    for doc in session.exec(select(Document).order_by(Document.uploaded_at.desc())).all():
        for n, page in enumerate(pages_of(doc), start=1):
            low = page.lower()
            score = sum(low.count(w) for w in words)
            if not score:
                continue
            at = min((low.find(w) for w in words if w in low), default=0)
            snippet = re.sub(r"\s+", " ", page[max(0, at - 160): at + 240]).strip()
            hits.append({"document_id": doc.id, "title": doc.title, "kind": doc.kind, "attach_to": doc.attach_to, "page": n, "score": score, "snippet": snippet})
    hits.sort(key=lambda h: -h["score"])
    return hits[:limit]


MONEY = r"\$ ?(\d{1,3}(?:,\d{3})+(?:\.\d{2})?|\d+(?:\.\d{2})?)(?![\d,])"
TERMS = {
    "rent": re.compile(r"(?<!pet )(?<!pet\t)\b(?:monthly[ \t]{1,3})?rent(?:[ \t]{1,3}amount)?[ \t]{0,3}(?:of|is|:)?[ \t]{0,3}" + MONEY, re.IGNORECASE),
    "pet_rent": re.compile(r"pet\s+rent[^\n$]{0,30}" + MONEY, re.IGNORECASE),
    "deposit": re.compile(r"security\s+deposit[^\n$]{0,40}" + MONEY, re.IGNORECASE),
    "late_fee": re.compile(r"late\s+(?:fee|charge)[^\n$]{0,60}" + MONEY, re.IGNORECASE),
    "notice_days": re.compile(r"(\d{2,3})[ \t-]{0,3}days?['’]?[ \t]{1,3}(?:written[ \t]{1,3})?notice", re.IGNORECASE),
    "due_day": re.compile(r"(?:rent[ \t]{1,3}is[ \t]{1,3})?due[ \t]{1,3}on[ \t]{1,3}(?:or[ \t]{1,3}before[ \t]{1,3})?the[ \t]{1,3}(\d{1,2})(?:st|nd|rd|th)", re.IGNORECASE),
    "grace_days": re.compile(r"(?:grace[ \t]{1,3}period[ \t]{1,3}of[ \t]{1,3}|after[ \t]{1,3}the[ \t]{1,3})(\d{1,2})(?:[ \t]{1,3}days|(?:st|nd|rd|th))", re.IGNORECASE),
    "ends_on": re.compile(r"(?:lease[ \t]{1,3}(?:term[ \t]{1,3})?(?:ends?|expires?|ending|expiration)(?:[ \t]{1,3}on)?|end[ \t]{1,3}date)[: \t]{1,4}([A-Za-z]{3,9}\.? \d{1,2},? \d{4}|\d{1,2}/\d{1,2}/\d{4})", re.IGNORECASE),
}


TERMS_SCAN_CHARS = 200_000  # a lease states its terms in its first pages; the rest is not scanned


def terms_of(pages: list[str]) -> dict[str, Any]:
    """What a lease or a charge sheet states, each with the page it came from. Deterministic; nothing guessed."""
    from ..mail.classify import parse_date_text

    found: dict[str, Any] = {}
    budget = TERMS_SCAN_CHARS
    for n, page in enumerate(pages, start=1):
        if budget <= 0:
            break
        text = page[:min(MAX_PAGE_CHARS, budget)]
        budget -= len(text)
        for key, pattern in TERMS.items():
            if key in found:
                continue
            m = pattern.search(text)
            if not m:
                continue
            raw = m.group(1)
            if key in ("rent", "pet_rent", "deposit", "late_fee"):
                value: Any = round(float(raw.replace(",", "")), 2)
            elif key in ("notice_days", "due_day", "grace_days"):
                value = int(raw)
                if (key == "due_day" and not 1 <= value <= 31) or (key == "notice_days" and not 1 <= value <= 180):
                    continue
                if key == "grace_days" and "after" in m.group(0).lower():
                    due = int((found.get("due_day") or {}).get("value") or 1)
                    value = max(0, value - due)
            else:
                when = parse_date_text(raw)
                if when is None:
                    continue
                value = when.isoformat()
            found[key] = {"value": value, "page": n, "quote": re.sub(r"\s+", " ", m.group(0))[:160]}
    return found


def offer_terms(session: Session, doc: Document, pages: list[str]) -> ExtractedFact | None:
    """The terms found, as one fact to confirm on Needs you; confirming applies them to the home."""
    terms = terms_of(pages)
    if not terms:
        return None
    fact = ExtractedFact(kind="document_terms", fields=json.dumps({"document_id": doc.id, "title": doc.title, "attach_to": doc.attach_to, "terms": terms}),
                         confidence=0.7, status="pending", excerpt="; ".join(f"{k}: {v['quote']}" for k, v in list(terms.items())[:4])[:600], produced_by="documents")
    session.add(fact)
    session.commit()
    return fact


def apply_terms(session: Session, fields: dict[str, Any], choose: dict[str, Any] | None = None) -> str:
    """Confirmed terms onto the home they belong to. `choose` (from the owner) leaves a term out (None) or gives it a
    corrected value; terms it does not name apply as read. Returns what changed, for the record."""
    from datetime import date as _date

    from ..models import Home, HomeCharge
    from .housing import charges_of, sync_bills

    attach = str(fields.get("attach_to") or "")
    if not attach.startswith("home:"):
        return ""
    home = session.get(Home, int(attach.split(":")[1]))
    if home is None:
        return ""
    terms = dict(fields.get("terms") or {})
    for key, value in (choose or {}).items():
        if key not in TERMS:
            continue
        if value is None:
            terms.pop(key, None)
        else:
            terms[key] = {**(terms.get(key) or {}), "value": value}
    changed = []
    for key in ("due_day", "grace_days", "notice_days"):
        if key in terms:
            setattr(home, key, int(terms[key]["value"]))
            changed.append(key)
    if "late_fee" in terms:
        home.late_fee, _ = float(terms["late_fee"]["value"]), changed.append("late_fee")
    if "deposit" in terms:
        home.deposit, _ = float(terms["deposit"]["value"]), changed.append("deposit")
    if "ends_on" in terms:  # the lease term's end; moving out is the owner's own date
        home.lease_ends_on, _ = _date.fromisoformat(terms["ends_on"]["value"]), changed.append("lease_ends_on")
    charges = charges_of(session, home.id)
    for key, name, kind in (("rent", "Rent", "rent"), ("pet_rent", "Pet rent", "pet")):
        if key in terms:
            c = next((c for c in charges if c.kind == kind), None) or HomeCharge(home_id=home.id, name=name, kind=kind)
            c.amount = float(terms[key]["value"])
            session.add(c)
            changed.append(key)
    session.add(home)
    session.commit()
    sync_bills(session, home)
    return f"home:{home.id}:" + ",".join(changed)
